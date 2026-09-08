from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "WashNgo_App_Features.docx"

NAVY = "17324D"
BLUE = "2F6BFF"
PALE = "EAF2FF"
PALE2 = "F5F8FC"
GRAY = "5F6B76"
LIGHT = "D9E1E8"
GREEN = "15805A"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=110, start=140, bottom=110, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_cell_border(cell, color=LIGHT):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}"); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "5"); node.set(qn("w:color"), color)

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_feature_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.55), Inches(2.7), Inches(2.75)]
    headers = ["Feature area", "What users can do", "Value to the workflow"]
    for i, text in enumerate(headers):
        c = table.rows[0].cells[i]; c.width = widths[i]; c.text = text
        set_cell_shading(c, NAVY); set_cell_margins(c); set_cell_border(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255); r.font.size = Pt(9.5)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = widths[i]; cells[i].text = text
            set_cell_margins(cells[i]); set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2: set_cell_shading(cells[i], PALE2)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
                for r in p.runs: r.font.size = Pt(9.4)
        cells[0].paragraphs[0].runs[0].font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_tagline(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10.7); styles["Normal"].font.color.rgb = RGBColor.from_string("24313D")
styles["Normal"].paragraph_format.space_after = Pt(7); styles["Normal"].paragraph_format.line_spacing = 1.13
for name, size, color in (("Title", 30, "000000"), ("Heading 1", 19, "000000"), ("Heading 2", 13, "000000")):
    s = styles[name]; s.font.name = "Aptos Display"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(14 if name != "Title" else 0); s.paragraph_format.space_after = Pt(7)
if "Feature Label" not in styles:
    s = styles.add_style("Feature Label", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Aptos"; s.font.size = Pt(9); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(BLUE)
    s.paragraph_format.space_after = Pt(5)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(50)
logo = ROOT / "assets" / "images" / "icon.png"
if logo.exists():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(logo), width=Inches(1.15))
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("WashNgo App Features")
add_tagline(doc, "Product feature overview for the laundry pickup and delivery MVP")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Lipa City Batangas  |  September 2026")
r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(GRAY)
doc.add_paragraph().paragraph_format.space_after = Pt(26)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("WashNgo connects customers, laundry partners, riders, and platform administrators in one role-based workflow—from booking to pickup, cleaning, return delivery, and completion.")
r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(NAVY)
p.paragraph_format.left_indent = Inches(.55); p.paragraph_format.right_indent = Inches(.55); p.paragraph_format.line_spacing = 1.25
doc.add_paragraph().paragraph_format.space_after = Pt(45)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prepared from the implemented Expo and Supabase application")
r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(GRAY)

doc.add_page_break()

doc.add_heading("Product Overview", level=1)
doc.add_paragraph("WashNgo is a multi-sided laundry marketplace designed for local pickup and delivery. The current MVP brings the entire order lifecycle into a single system, giving each role only the screens and actions relevant to their work.")
doc.add_heading("Who the app serves", level=2)
add_feature_table(doc, [
    ("Customers", "Find verified laundries, select a service, schedule pickup, view pricing, and track orders.", "Removes the need to travel, call shops, or manually follow up on an order."),
    ("Laundry partners", "Review bookings, accept or reject requests, update cleaning progress, and monitor revenue.", "Creates a simple operating queue and a consistent handoff process."),
    ("Riders", "Go online, accept eligible jobs, update job status, and submit photo and location proofs.", "Makes pickup and return delivery traceable and accountable."),
    ("Administrators", "View platform metrics and approve or reject rider and partner verification requests.", "Supports controlled onboarding and launch operations."),
])
doc.add_heading("Core product promise", level=2)
doc.add_paragraph("A customer can book a verified laundry partner in minutes, see an itemized estimate, and follow every handoff. Laundry partners and riders receive controlled action queues, while administrators retain oversight of account verification and platform activity.")
doc.add_heading("At a glance", level=2)
for t in [
    "Role-based access for customer, rider, laundry partner, and administrator accounts",
    "Two-leg logistics with separate customer-to-laundry and laundry-to-customer delivery jobs",
    "Realtime order status and immutable handoff history",
    "Server-calculated pricing with laundry, delivery, and platform fee breakdowns",
    "Private proof storage, database authorization policies, and transactional lifecycle controls",
]: add_bullet(doc, t)

doc.add_page_break()
doc.add_heading("Customer Features", level=1)
doc.add_paragraph("The customer experience covers account access, discovery, booking, and order visibility. It is optimized for a straightforward path from selecting a laundry service to receiving clean clothes.")
add_feature_table(doc, [
    ("Account access", "Create a customer account, sign in, stay signed in through a persisted session, reset a forgotten password, and sign out.", "Reduces sign-in friction while keeping each user in the correct role-based area."),
    ("Saved addresses", "Add pickup addresses with label, barangay, full address, and map coordinates; automatically mark the first address as default.", "Speeds up repeat bookings and gives delivery jobs precise locations."),
    ("Laundry discovery", "Browse verified and active Lipa City partners, search by shop name, see ratings, addresses, service names, and starting prices.", "Helps customers compare available partners before booking."),
    ("Service selection", "View service descriptions, per-kilogram or per-load prices, minimum charges, and estimated turnaround time.", "Sets clear expectations for cost and service duration."),
    ("Pickup booking", "Choose an address, date, time, estimated weight, and special instructions before confirming.", "Captures the information needed for fulfillment in one guided flow."),
    ("Transparent estimate", "See the stored total and its laundry subtotal, pickup fee, return fee, and platform fee components.", "Makes the cost structure understandable before fulfillment progresses."),
    ("Order center", "View an active-order shortcut, recent orders, complete order history, status badges, service details, and total amount.", "Keeps current and previous transactions easy to find."),
    ("Live tracking", "View the pickup and laundry points on a map plus a realtime status and timestamped handoff timeline.", "Reduces uncertainty and provides an auditable order history."),
])
doc.add_heading("Customer payment model", level=2)
doc.add_paragraph("The MVP supports cash on delivery and pay later as internal ledger states. It tracks amount due, amount paid, remaining balance, and payment status. A live payment gateway is not yet included.")

doc.add_page_break()
doc.add_heading("Laundry Partner Features", level=1)
doc.add_paragraph("The partner workspace presents the actions a laundry shop needs to move orders safely through confirmation and cleaning.")
add_feature_table(doc, [
    ("Order queue", "View active bookings with order number, pickup schedule, estimated weight, status, and total.", "Gives staff one operational list for incoming and in-progress work."),
    ("Booking decision", "Accept a booking to begin fulfillment or reject it with a recorded reason.", "Prevents unconfirmed orders from entering dispatch."),
    ("Cleaning workflow", "Mark received orders as processing, then mark them ready for return.", "Keeps customer status and delivery dispatch synchronized with actual shop progress."),
    ("Operations metrics", "See active order count and completed-order laundry revenue.", "Provides an immediate snapshot of workload and earned service revenue."),
    ("Account status", "View profile and verification state from the partner workspace.", "Makes platform eligibility and account identity visible."),
])
doc.add_heading("Rider Features", level=1)
doc.add_paragraph("The rider dashboard separates availability, job acceptance, execution, proof collection, and earnings.")
add_feature_table(doc, [
    ("Approval gate", "Access delivery jobs only after administrator approval.", "Limits dispatch to verified rider accounts."),
    ("Online availability", "Go online or offline; going online records the rider's current foreground location.", "Controls when a rider is eligible to receive nearby jobs."),
    ("Job marketplace", "See eligible pickup or return jobs, their direction, and rider payout; accept an available job atomically.", "Avoids double assignment when multiple riders act at once."),
    ("Guided delivery", "Move through accepted, arriving, picked up, and completed states.", "Standardizes the handoff sequence for both delivery legs."),
    ("Proof of handoff", "Capture a camera photo and foreground GPS location for pickup and delivery completion.", "Creates traceable evidence before sensitive status transitions are accepted."),
    ("Earnings summary", "See completed job count and total payout from completed work.", "Gives riders a concise performance and earnings view."),
])

doc.add_page_break()
doc.add_heading("Admin and Platform Features", level=1)
doc.add_heading("Administrator controls", level=2)
add_feature_table(doc, [
    ("Launch dashboard", "View total users, active orders, orders created today, and platform revenue.", "Supports daily operational monitoring."),
    ("Verification queue", "Review pending rider and laundry partner accounts and approve or reject each request.", "Centralizes controlled access to supply-side roles."),
    ("Responsive access", "Use the admin experience through the responsive Expo Web interface.", "Allows launch operations without a separate admin application."),
])
doc.add_heading("Shared platform capabilities", level=2)
add_feature_table(doc, [
    ("Role protection", "Restore the signed-in session and route users to the area tied to their database-owned role.", "Prevents client-selected role escalation."),
    ("Realtime updates", "Refresh orders and status logs when relevant database changes occur.", "Keeps customer tracking and operational views current."),
    ("Notifications", "Create in-app order notifications and support Expo push delivery through a secured Edge Function webhook.", "Allows lifecycle events to reach users beyond the active screen."),
    ("Transactional actions", "Use database functions for booking, job acceptance, lifecycle transitions, location updates, and verification.", "Keeps multi-record changes consistent and authorized."),
    ("Service zones", "Use PostGIS-backed zones to associate shops, riders, and eligible jobs with the launch area.", "Provides a foundation for geographic dispatch."),
    ("Data security", "Apply Row Level Security to app data and owner-scoped policies to private proof and verification files.", "Restricts data to authorized participants and administrators."),
    ("Auditability", "Record immutable order status logs and retain delivery proofs with timestamps and coordinates.", "Supports dispute review and operational accountability."),
])

doc.add_page_break()
doc.add_heading("End to End Order Flow", level=1)
doc.add_paragraph("WashNgo models pickup and return as separate delivery jobs. This prevents one generic delivery state from hiding where an order is in the physical workflow.")
steps = [
    ("1", "Customer booking", "The customer selects a verified shop and service, chooses a saved address and schedule, enters estimated weight and instructions, and confirms the request."),
    ("2", "Partner confirmation", "The laundry partner accepts or rejects the booking. Acceptance creates the pickup dispatch stage."),
    ("3", "Pickup delivery", "An approved online rider accepts the customer-to-laundry job, confirms arrival and pickup, provides proof, and completes delivery to the shop."),
    ("4", "Laundry processing", "The partner confirms receipt, starts processing, and marks the order ready for return."),
    ("5", "Return delivery", "A rider accepts the laundry-to-customer job, completes the return handoff with proof, and the order reaches delivered status."),
    ("6", "Completion", "The completed transaction remains available in order history with totals, statuses, and its timestamped handoff log."),
]
table = doc.add_table(rows=0, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
for n, name, detail in steps:
    cells = table.add_row().cells
    for i,w in enumerate((Inches(.55), Inches(1.65), Inches(4.8))): cells[i].width=w
    cells[0].text=n; cells[1].text=name; cells[2].text=detail
    for c in cells:
        set_cell_margins(c, 130, 140, 130, 140); set_cell_border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cells[0], BLUE)
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cells[0].paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    cells[1].paragraphs[0].runs[0].font.bold=True
    for c in cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
            for r in p.runs: r.font.size=Pt(9.4)
doc.add_paragraph()
doc.add_page_break()
doc.add_heading("Current MVP boundaries", level=1)
doc.add_paragraph("The following items are intentionally outside the current interface or require production configuration. They are not represented as fully available end-user features.")
for t in [
    "No integrated PayMongo, GCash, Maya, or card gateway; payments remain ledger records.",
    "Dispatch is an eligible-zone broadcast with atomic acceptance, not route optimization.",
    "Rider location uses foreground updates when going online; background tracking is not implemented.",
    "Push delivery requires Expo credentials plus the configured database webhook and Edge Function secret.",
    "Ratings are supported by the data and service layer, but the current customer screens do not yet expose a rating form.",
    "Chat, OTP handoff, and QR verification have schema preparation but no current user interface.",
]: add_bullet(doc, t)
doc.add_heading("Recommended next releases", level=2)
doc.add_paragraph("Prioritize integrated digital payments, background rider tracking with clear privacy controls, optimized dispatch, customer support chat, configurable fees, partner calendars, refunds, promotions, and multi-city operations.")

# Headers and footers
for section in doc.sections:
    hp = section.header.paragraphs[0]; hp.text = "WashNgo  |  App Features"; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("WashNgo Product Documentation  •  ")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); fp._p.append(fld)
    for r in fp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
